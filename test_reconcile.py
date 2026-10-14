import io
import unittest
import fitz
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from reconcile import compare, extract_pdf, suggest_rows, report_bytes, money


def frame(rows):
    return pd.DataFrame(rows, columns=['Дата','Документ','Сумма'])

class ReconciliationTests(unittest.TestCase):
    def test_amounts_and_missing(self):
        a=frame([['01.01.2026','1','1 234,56'],['02.01.2026','2','200,00'],['03.01.2026','3','50,00']])
        b=frame([['01.01.2026','1','1234.56'],['02.01.2026','2','210,00'],['04.01.2026','4','75,00']])
        r=compare(a,b)
        self.assertEqual(list(r['Статус']),['Совпадает','Расхождение','Только у нас','Только у контрагента'])
        self.assertEqual(r.iloc[1]['Разница'],-10)
    def test_duplicates_not_silently_matched(self):
        a=frame([['01.01.2026','1','100']]*2)
        b=frame([['01.01.2026','1','100']])
        self.assertEqual(list(compare(a,b)['Статус']),['Неоднозначное совпадение']*2+['Только у контрагента'])
    def test_dates_signs_and_invalid(self):
        a=frame([['01.01.2026','1','100']]); b=frame([['02.01.2026','1','100']])
        self.assertEqual(compare(a,b).iloc[0]['Статус'],'Разница дат')
        b.loc[0,'Сумма']='-100'
        self.assertEqual(compare(a,b).iloc[0]['Статус'],'Расхождение')
        with self.assertRaises(ValueError): money('nan')
        a.loc[0,'Дата']='invalid'
        with self.assertRaises(ValueError): compare(a,b)
    def test_report_formula_safety(self):
        a=frame([['01.01.2026','=1+1','100']])
        output=report_bytes(compare(a,a))
        ws=load_workbook(io.BytesIO(output)).active
        self.assertEqual(ws['B2'].data_type,'s')
        self.assertEqual(ws['B2'].value,'=1+1')
    def test_real_scan_ocr(self):
        image=Image.new('RGB',(1800,600),'white')
        draw=ImageDraw.Draw(image)
        font=ImageFont.truetype('/usr/share/fonts/truetype/open-sans/OpenSans-Regular.ttf',48)
        draw.text((60,60),'Акт сверки',font=font,fill='black')
        draw.text((60,160),'01.01.2026 Документ № 123   1 234,56',font=font,fill='black')
        buf=io.BytesIO(); image.save(buf,format='PNG')
        pdf=fitz.open(); page=pdf.new_page(width=900,height=300)
        page.insert_image(page.rect,stream=buf.getvalue())
        text=extract_pdf(pdf.tobytes()); pdf.close()
        self.assertIn('01.01.2026',text)
        rows=suggest_rows(text)
        self.assertEqual(len(rows),1)
        self.assertEqual(money(rows.iloc[0]['Сумма']),123456)
        self.assertEqual(rows.iloc[0]['Документ'],'123')

if __name__=='__main__': unittest.main()
